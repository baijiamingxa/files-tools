#!/usr/bin/env python3
"""文档转换与安装器离线保护测试；--real 运行真实转换与渲染。"""
from __future__ import annotations

import json
import os
import tempfile
import threading
import time
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

import document_conversion as core
import document_runtime as engine
from file_tools import FileToolCancelled
import installer_core as installer
from txt_encoder_core import make_console_tolerant


def docx_fixture(path):
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("word/document.xml", '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body/></w:document>')
        archive.writestr("[Content_Types].xml", '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>')


def fake_worker(job, workdir, *args):
    if job["mode"] == "pdf_word":
        docx_fixture(job["output"])
    else:
        Path(job["output"]).write_bytes(b"%PDF-1.4\nfixture\n%%EOF")
    return {"ok": True, "detail": "模拟转换完成"}


class DocumentTests(unittest.TestCase):
    def test_plan_routes_and_rejects_collisions(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            docx_fixture(root / "中文 &.docx")
            (root / "中文 &.doc").write_bytes(b"old word")
            (root / "other.pdf").write_bytes(b"%PDF-fixture")
            rows = core.build_document_plan(list(root.iterdir()), core.DocumentOptions())
            self.assertEqual(sum(bool(row.error) for row in rows), 2)
            self.assertTrue(next(row for row in rows if row.path.suffix == ".pdf").destination.name.endswith(".docx"))
            rows = core.build_document_plan([root / "other.pdf"], core.DocumentOptions(mode="word_pdf"))
            self.assertEqual(rows[0].status, "已跳过")

    def test_source_preserved_and_existing_output_not_overwritten(self):
        with tempfile.TemporaryDirectory() as temp, patch.object(core, "run_document_worker", side_effect=fake_worker):
            root = Path(temp)
            source = root / "源文件.pdf"
            source.write_bytes(b"%PDF-original")
            options = core.DocumentOptions()
            rows = core.build_document_plan([source], options)
            core.execute_document_plan(rows, options)
            self.assertEqual(rows[0].status, "已完成", rows[0].error)
            self.assertEqual(source.read_bytes(), b"%PDF-original")
            original_output = rows[0].destination.read_bytes()
            blocked = core.build_document_plan([source], options)
            self.assertTrue(blocked[0].error)
            self.assertEqual(rows[0].destination.read_bytes(), original_output)

    def test_stale_source_and_publish_race(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "sample.pdf"
            source.write_bytes(b"%PDF-first")
            options = core.DocumentOptions()
            rows = core.build_document_plan([source], options)
            stat = source.stat()
            source.write_bytes(b"%PDF-other")
            os.utime(source, ns=(stat.st_atime_ns, stat.st_mtime_ns))
            core.execute_document_plan(rows, options)
            self.assertEqual(rows[0].status, "失败")
            self.assertFalse(rows[0].destination.exists())
            rows = core.build_document_plan([source], options)
            def race(job, workdir, *args):
                result = fake_worker(job, workdir, *args)
                rows[0].destination.write_bytes(b"other process")
                return result
            with patch.object(core, "run_document_worker", side_effect=race):
                core.execute_document_plan(rows, options)
            self.assertEqual(rows[0].status, "失败")
            self.assertEqual(rows[0].destination.read_bytes(), b"other process")

    def test_engine_failure_does_not_block_pdf_and_bad_results_rejected(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "word.docx"
            docx_fixture(source)
            pdf = root / "pdf.pdf"
            pdf.write_bytes(b"%PDF-fixture")
            options = core.DocumentOptions()
            rows = core.build_document_plan([source, pdf], options)
            with patch.object(engine, "ensure_office", side_effect=OSError("offline")), patch.object(core, "run_document_worker", side_effect=fake_worker):
                core.execute_document_plan(rows, options)
            self.assertEqual([row.status for row in rows], ["失败", "已完成"])
            options = core.DocumentOptions(output_dir=str(root / "bad"))
            rows = core.build_document_plan([pdf], options)
            def bad(job, *args):
                Path(job["output"]).write_bytes(b"not docx")
                return {"ok": True}
            with patch.object(core, "run_document_worker", side_effect=bad):
                core.execute_document_plan(rows, options)
            self.assertEqual(rows[0].status, "失败")
            self.assertFalse(rows[0].destination.exists())

    def test_cancel_stops_owned_child_processes(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            worker = root / "fixture_worker.py"
            worker.write_text("""import json,sys,time,subprocess
from pathlib import Path
j=json.loads(Path(sys.argv[1]).read_text(encoding='utf-8'))
while not Path(j['gate']).exists():time.sleep(.02)
p=Path(sys.argv[1]).parent
child='import time;from pathlib import Path;time.sleep(1.5);Path('+repr(str(p/'escaped.txt'))+').write_text(\"bad\")'
subprocess.Popen([sys.executable,'-c',child])
(p/'child-started').touch()
time.sleep(20)
""", encoding="utf-8")
            stop = threading.Event()
            def cancel():
                deadline = time.monotonic() + 10
                while not (root / "child-started").exists() and time.monotonic() < deadline:
                    time.sleep(0.02)
                stop.set()
            thread = threading.Thread(target=cancel)
            thread.start()
            with patch.object(core, "_worker_command", side_effect=lambda path: [__import__('sys').executable, str(worker), str(path)]):
                with self.assertRaises(FileToolCancelled):
                    core.run_document_worker({}, root, stop, timeout=12)
            thread.join()
            time.sleep(1.7)
            self.assertFalse((root / "escaped.txt").exists(), "子进程不能在取消后继续运行")

    def test_msi_directory_paths_are_contained(self):
        records = [("ROOT", "", "SourceDir"), ("INSTALLLOCATION", "ROOT", "LibreOffice"), ("PROGRAM", "INSTALLLOCATION", "program"), ("DATA", "PROGRAM", "SHORT|Long Name")]
        result = engine.directory_map(records)
        self.assertEqual(result["DATA"], Path("program") / "Long Name")
        with self.assertRaises(core.FileToolError):
            engine.directory_map([*records, ("BAD", "INSTALLLOCATION", "../escape")])


class InstallerTests(unittest.TestCase):
    def test_install_upgrade_preserves_unlisted_files(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "source"
            source.mkdir()
            (source / "文件工具箱.exe").write_bytes(b"version1")
            (source / "runtime.dll").write_bytes(b"library")
            package, target = root / "app.zip", root / "installed"
            installer.create_payload(source, package, "1.2.0")
            installer.install_payload(package, target)
            self.assertEqual((target / "文件工具箱.exe").read_bytes(), b"version1")
            (target / "个人文件.txt").write_bytes(b"keep")
            (source / "文件工具箱.exe").write_bytes(b"version2")
            installer.create_payload(source, package, "1.2.1")
            installer.install_payload(package, target)
            self.assertEqual((target / "文件工具箱.exe").read_bytes(), b"version2")
            self.assertEqual((target / "个人文件.txt").read_bytes(), b"keep")

    def test_install_rejects_unsafe_destination_and_corrupt_payload(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "source"
            source.mkdir()
            (source / "文件工具箱.exe").write_bytes(b"app")
            package = root / "app.zip"
            installer.create_payload(source, package, "1.2.0")
            with self.assertRaises(ValueError):
                installer.install_payload(package, source)
            bad = root / "bad.zip"
            with zipfile.ZipFile(bad, "w") as archive:
                archive.writestr("payload.json", json.dumps({"id": installer.INSTALL_ID, "files": {"../escape.exe": "0" * 64}}))
            with self.assertRaises(ValueError):
                installer.install_payload(bad, root / "new")
            with zipfile.ZipFile(package) as original, zipfile.ZipFile(bad, "w") as archive:
                archive.writestr("payload.json", original.read("payload.json"))
                archive.writestr("文件工具箱.exe", b"corrupt")
            with self.assertRaises(ValueError):
                installer.install_payload(bad, root / "new")
            self.assertFalse((root / "new" / "文件工具箱.exe").exists())

    def test_failed_upgrade_rolls_back(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "source"
            source.mkdir()
            for name in ("文件工具箱.exe", "library.dll"):
                (source / name).write_bytes(b"old")
            package, target = root / "app.zip", root / "installed"
            installer.create_payload(source, package, "1.2.0")
            installer.install_payload(package, target)
            old_manifest = (target / installer.MANIFEST).read_bytes()
            for path in source.iterdir():
                path.write_bytes(b"new")
            installer.create_payload(source, package, "1.2.1")
            def fail(done, total, text):
                if done == total:
                    raise OSError("simulated disk failure")
            with self.assertRaises(OSError):
                installer.install_payload(package, target, fail)
            self.assertEqual((target / "文件工具箱.exe").read_bytes(), b"old")
            self.assertEqual((target / "library.dll").read_bytes(), b"old")
            self.assertEqual((target / installer.MANIFEST).read_bytes(), old_manifest)


def real_test():
    from docx import Document
    from docx.shared import Inches
    import pymupdf
    if not engine.find_office():
        raise RuntimeError("真实测试需要已有 LibreOffice。")
    root = Path(__file__).resolve().parent / "_文档转换验证"
    root.mkdir(exist_ok=True)
    source = root / "中文 表格样本.docx"
    document = Document()
    document.sections[0].page_width = Inches(8.5)
    document.sections[0].page_height = Inches(11)
    document.add_heading("Word and PDF conversion", 0)
    document.add_paragraph("中文文字与 English content 123。文档转换验证。")
    table = document.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    for row, values in zip(table.rows, (("项目", "结果"), ("文字与表格", "保留"), ("源文件", "不覆盖"))):
        for cell, text in zip(row.cells, values):
            cell.text = text
    document.save(source)
    run_id = str(time.time_ns())
    options = core.DocumentOptions(output_dir=str(root / ("PDF-" + run_id)))
    rows = core.build_document_plan([source], options)
    core.execute_document_plan(rows, options)
    assert rows[0].status == "已完成", rows[0].error
    pdf_path = rows[0].destination
    with pymupdf.open(pdf_path) as pdf:
        assert "English content" in pdf[0].get_text()
        pdf[0].get_pixmap(matrix=pymupdf.Matrix(1.4, 1.4)).save(str(root / "Word转PDF.png"))
        scanned = pymupdf.open()
        for _ in range(2):
            page = scanned.new_page(width=pdf[0].rect.width, height=pdf[0].rect.height)
            page.insert_image(page.rect, stream=pdf[0].get_pixmap(matrix=pymupdf.Matrix(1.3, 1.3)).tobytes("png"))
        scanned_path = root / "扫描图片两页.pdf"
        scanned.save(scanned_path)
        scanned.close()
    options = core.DocumentOptions(output_dir=str(root / ("Word-" + run_id)))
    rows = core.build_document_plan([pdf_path, scanned_path], options)
    core.execute_document_plan(rows, options)
    assert all(row.status == "已完成" for row in rows), [(row.status, row.error) for row in rows]
    assert "English content" in "\n".join(paragraph.text for paragraph in Document(rows[0].destination).paragraphs)
    assert len(Document(rows[1].destination).inline_shapes) == 2
    assert "不可直接编辑" in rows[1].detail
    options = core.DocumentOptions(output_dir=str(root / ("roundtrip-" + run_id)))
    rendered = core.build_document_plan([row.destination for row in rows], options)
    core.execute_document_plan(rendered, options)
    assert all(row.status == "已完成" for row in rendered), [(row.status, row.error) for row in rendered]
    for index, row in enumerate(rendered):
        with pymupdf.open(row.destination) as pdf:
            if index == 1:
                assert pdf.page_count == 2, pdf.page_count
            pdf[0].get_pixmap(matrix=pymupdf.Matrix(1.4, 1.4)).save(str(root / f"Word往返预览{index}.png"))
    print(f"真实 Word / PDF 双向转换、中文表格和两页扫描件验证全部通过：{root}")


if __name__ == "__main__":
    make_console_tolerant()
    suite = unittest.defaultTestLoader.loadTestsFromModule(__import__(__name__))
    result = unittest.TextTestRunner(verbosity=2).run(suite)
    if not result.wasSuccessful():
        raise SystemExit(1)
    print("文档转换和安装器自检全部通过 ✔")
    if "--real" in __import__("sys").argv:
        real_test()
