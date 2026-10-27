#!/usr/bin/env python3
"""文档转换专用工作进程；不打开工具箱主窗口。"""
from __future__ import annotations

import contextlib
import io
import json
import subprocess
import sys
import time
from pathlib import Path


def _write(path, data):
    path = Path(path)
    temp = path.with_suffix(".tmp")
    temp.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    temp.replace(path)


def convert_pdf(job, report):
    import pymupdf
    from docx import Document
    from docx.enum.section import WD_SECTION_START
    from docx.shared import Pt
    from pdf2docx import Converter
    password = job.get("password", "")
    with pymupdf.open(job["source"]) as pdf:
        if pdf.needs_pass and not pdf.authenticate(password):
            raise ValueError("PDF 已加密，密码为空或不正确。")
        if not pdf.page_count:
            raise ValueError("PDF 没有可转换的页面。")
        text_pages = sum(bool(page.get_text("text").strip()) for page in pdf)
        count = pdf.page_count
        if not text_pages:
            report("扫描 / 图像 PDF：正在保留页面图片…")
            document = Document()
            normal = document.styles["Normal"]
            normal.font.size = Pt(1)
            normal.paragraph_format.space_before = Pt(0)
            normal.paragraph_format.space_after = Pt(0)
            for index, page in enumerate(pdf):
                section = document.sections[0] if index == 0 else document.add_section(WD_SECTION_START.NEW_PAGE)
                width, height = page.rect.width, page.rect.height
                section.page_width, section.page_height = Pt(width), Pt(height)
                section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Pt(12)
                scale = min((width - 24) / width, (height - 36) / height)
                image = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False).tobytes("png")
                paragraph = document.add_paragraph()
                paragraph.add_run().add_picture(io.BytesIO(image), width=Pt(width * scale), height=Pt(height * scale))
                report(f"图像页 {index + 1} / {count}")
            document.save(job["output"])
            return f"已转换 {count} 页；扫描 / 图像版 PDF 已保留为 Word 页面图片，未进行 OCR，图片内文字不可直接编辑。"
    report(f"正在重建 {count} 页文字、图片和表格…")
    converter = Converter(job["source"], password=password or None)
    try:
        converter.convert(job["output"], start=0, end=None, multi_processing=False)
    finally:
        converter.close()
    # 重新打开 OOXML，确认结果可读，而不只依赖子进程返回码。
    Document(job["output"])
    return f"已转换 {count} 页为可编辑 DOCX；复杂分页、字体和表格可能需要手工调整。"


def convert_word(job, report):
    import pymupdf
    office = Path(job["office"])
    if not office.is_file():
        raise ValueError("未找到 Word 转 PDF 引擎，请先配置 LibreOffice。")
    work = Path(job["output"]).parent
    profile = work / "office-profile"
    (profile / "user").mkdir(parents=True)
    (profile / "user" / "registrymodifications.xcu").write_text('''<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry"><item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item></oor:items>''', encoding="utf-8")
    report("正在使用独立配置转换 Word 文档…")
    command = [str(office), "-env:UserInstallation=" + profile.as_uri(), "--headless", "--norestore", "--nodefault", "--nologo",
               "--convert-to", "pdf:writer_pdf_Export", "--outdir", str(work), job["source"]]
    result = subprocess.run(command, stdin=subprocess.DEVNULL, stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
        creationflags=0x08000000 if sys.platform == "win32" else 0)
    generated = work / (Path(job["source"]).stem + ".pdf")
    if result.returncode != 0 or not generated.is_file():
        raise ValueError("Word 转 PDF 失败，文档可能受密码保护、损坏或格式不受支持。" + result.stdout.decode("utf-8", "replace")[-500:])
    with pymupdf.open(generated) as pdf:
        pages = pdf.page_count
        if pages == 0:
            raise ValueError("转换结果没有页面。")
    generated.replace(job["output"])
    return f"已转换为 {pages} 页 PDF（LibreOffice）；源文档保留。"


def open_pdf(path, password=""):
    import pymupdf
    pdf = pymupdf.open(path)
    try:
        if not pdf.is_pdf:
            raise ValueError("文件内容不是 PDF。")
        if pdf.needs_pass and not pdf.authenticate(password):
            raise ValueError("PDF 已加密，密码为空或不正确。")
        if not pdf.page_count:
            raise ValueError("PDF 没有可用页面。")
        return pdf
    except Exception:
        pdf.close()
        raise


def edit_pdf(job, report):
    import pymupdf
    from contextlib import ExitStack
    with ExitStack() as stack:
        sources = [stack.enter_context(open_pdf(path, job.get("password", ""))) for path in job["sources"]]
        operation, outputs = job["operation"], job["outputs"]
        expected = []
        if operation == "merge":
            with pymupdf.open() as target:
                for index, source in enumerate(sources, 1):
                    report(f"合并文件 {index} / {len(sources)}")
                    target.insert_pdf(source)
                expected.append(target.page_count)
                target.save(outputs[0], garbage=3, deflate=True)
        elif operation == "split":
            for index, (selection, path) in enumerate(zip(job["selections"], outputs), 1):
                report(f"生成分卷 {index} / {len(outputs)}")
                with pymupdf.open() as target:
                    target.insert_pdf(sources[0], from_page=selection[0], to_page=selection[-1])
                    expected.append(target.page_count)
                    target.save(path, garbage=3, deflate=True)
        elif operation in ("extract", "rotate"):
            source = sources[0]
            if operation == "extract":
                source.select(job["selections"][0])
            else:
                for number in job["selections"][0]:
                    page = source[number]
                    page.set_rotation((page.rotation + job["rotation"]) % 360)
            expected.append(source.page_count)
            source.save(outputs[0], garbage=3, deflate=True, encryption=pymupdf.PDF_ENCRYPT_NONE)
        else:
            raise ValueError("未知 PDF 操作。")
    if len(outputs) != len(expected):
        raise ValueError("输出文件数量与计划不一致。")
    for path, count in zip(outputs, expected):
        with open_pdf(path) as pdf:
            if pdf.page_count != count:
                raise ValueError("生成 PDF 页数校验失败。")
    return f"已生成 {len(outputs)} 个 PDF，共 {sum(expected)} 页；源文件保留。"


def run_worker(path):
    job_path = Path(path)
    job = json.loads(job_path.read_text(encoding="utf-8"))
    deadline = time.monotonic() + 20
    while not Path(job["gate"]).exists():
        if time.monotonic() > deadline:
            return 2
        time.sleep(0.05)
    report = lambda message: _write(job["progress"], {"message": message})
    try:
        with open(job_path.parent / "engine.log", "w", encoding="utf-8") as log, contextlib.redirect_stdout(log), contextlib.redirect_stderr(log):
            if job["mode"] == "pdf_info":
                with open_pdf(job["source"], job.get("password", "")) as pdf:
                    _write(job["result"], {"ok": True, "pages": pdf.page_count})
                return 0
            if job["mode"] == "pdf_edit":
                detail = edit_pdf(job, report)
            elif job["mode"] == "pdf_word":
                detail = convert_pdf(job, report)
            elif job["mode"] == "word_pdf":
                detail = convert_word(job, report)
            else:
                raise ValueError("未知文档转换方向。")
        _write(job["result"], {"ok": True, "detail": detail})
        return 0
    except Exception as exc:
        _write(job["result"], {"ok": False, "error": str(exc)})
        return 1


if __name__ == "__main__":
    raise SystemExit(run_worker(sys.argv[1]))
